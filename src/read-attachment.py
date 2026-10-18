"""Read-only attachment extraction. Input JSON/base64; no remote fetching or OCR."""
import sys, json, base64, io, zipfile

request = json.load(sys.stdin)
content = base64.b64decode(request['contentBase64'], validate=True)
kind = request['kind']
segments = []
if kind in ('docx', 'xlsx'):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError('Decompressed attachment too large')
if kind == 'pdf':
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    if len(reader.pages) > 500:
        raise ValueError('PDF too many pages')
    segments = [{'locator': f'page:{index+1}', 'text': page.extract_text() or ''} for index, page in enumerate(reader.pages)]
elif kind == 'docx':
    from docx import Document
    document = Document(io.BytesIO(content))
    segments = [{'locator': f'paragraph:{index+1}', 'text': paragraph.text} for index, paragraph in enumerate(document.paragraphs)]
    for index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            segments.append({'locator': f'table:{index+1}/row:{row_index+1}', 'text': '\t'.join(cell.text for cell in row.cells)})
elif kind == 'xls':
    import xlrd
    workbook = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True)
    if workbook.nsheets > 500:
        raise ValueError('XLS too many sheets')
    total_cells = 0
    def address(row, column):
        letters = ''
        number = column + 1
        while number:
            number, remainder = divmod(number - 1, 26)
            letters = chr(65 + remainder) + letters
        return f'{letters}{row + 1}'
    try:
        for sheet in workbook.sheets():
            total_cells += sheet.nrows * sheet.ncols
            if total_cells > 500000:
                raise ValueError('XLS too many cells')
            merged = {}
            for first_row, last_row, first_column, last_column in sheet.merged_cells:
                if first_row < 0 or first_column < 0 or last_row > sheet.nrows or last_column > sheet.ncols or (last_row-first_row)*(last_column-first_column) > 500000:
                    raise ValueError('XLS merged range outside bounded sheet')
                anchor = address(first_row, first_column)
                span = f'{anchor}:{address(last_row - 1, last_column - 1)}'
                for row in range(first_row, last_row):
                    for column in range(first_column, last_column):
                        merged[(row, column)] = (first_row, first_column, anchor, span)
            for row in range(sheet.nrows):
                values, cells = [], []
                for column in range(sheet.ncols):
                    merge = merged.get((row, column))
                    source_row, source_column = merge[:2] if merge else (row, column)
                    cell = sheet.cell(source_row, source_column)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        text = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                    elif value == '':
                        text = ''
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                        text = str(int(value))
                    else:
                        text = str(value)
                    values.append(text)
                    if text or merge:
                        item = {'address': address(row, column), 'text': text, 'rawValue': value}
                        if merge:
                            item.update(anchor=merge[2], mergedRange=merge[3])
                        cells.append(item)
                segments.append({'locator': f'sheet:{sheet.name}/row:{row + 1}', 'text': '\t'.join(values).rstrip('\t'), 'cells': cells})
    finally:
        workbook.release_resources()
elif kind == 'xlsx':
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    for sheet in workbook:
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            segments.append({'locator': f'sheet:{sheet.title}/row:{index+1}', 'text': '\t'.join(str(value) if value is not None else '' for value in row)})
    workbook.close()
else:
    raise ValueError('Unsupported attachment type')
if not any(segment['text'].strip() for segment in segments):
    raise ValueError('No extractable text; scanned/encrypted attachments need manual text with source retained')
json.dump({'segments': segments, 'text': '\n'.join(segment['text'] for segment in segments)}, sys.stdout, ensure_ascii=False)
